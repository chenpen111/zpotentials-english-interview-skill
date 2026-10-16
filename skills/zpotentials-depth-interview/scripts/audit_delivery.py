"""Read-only structural checks; review evidence is not a substitute for review."""

import argparse
import hashlib
import json
import re
import sys
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from docx import Document
from skill_config import load_rules, render_footer

SKILL = Path(__file__).resolve().parents[1]
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
TIMESTAMP = re.compile(r"\b\d{1,2}:\d{2}(?::\d{2})?\b")


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def visible_text(doc):
    # XML includes hyperlink runs and table paragraphs omitted by plain paragraph.text.
    paragraphs = []
    for p in doc._element.body.xpath(".//w:p"):
        text = "".join(node.text or "" for node in p.xpath(".//w:t"))
        if text.strip():
            paragraphs.append(text)
    return "\n".join(paragraphs)


def comment_map(path):
    with ZipFile(path) as zf:
        if "word/comments.xml" not in zf.namelist():
            return {}, set()
        root = ET.fromstring(zf.read("word/comments.xml"))
        comments = {
            node.attrib[f"{{{NS['w']}}}id"]: "".join(node.itertext())
            for node in root.findall("w:comment", NS)
        }
        document = ET.fromstring(zf.read("word/document.xml"))
        anchors = {
            node.attrib[f"{{{NS['w']}}}id"]
            for node in document.findall(".//w:commentRangeStart", NS)
        }
    return comments, anchors


def image_roles(doc):
    counts = {"cover": 0, "body": 0, "unclassified": 0}
    for shape in doc.inline_shapes:
        props = shape._inline.docPr
        labels = " ".join(props.get(key, "") for key in ("name", "title", "descr"))
        if "ZP_COVER" in labels:
            counts["cover"] += 1
        elif "ZP_BODY" in labels:
            counts["body"] += 1
        else:
            counts["unclassified"] += 1
    return counts


def has_descriptive_caption(paragraphs, image_index):
    for paragraph in paragraphs[image_index + 1:]:
        if paragraph._p.xpath(".//w:drawing"):
            return False
        text = paragraph.text.strip()
        if text:
            return (not text.startswith(("图片来源", "图源"))
                    and not re.search(r"https?://", text)
                    and not paragraph.style.name.startswith("Heading")
                    and not any(run.bold for run in paragraph.runs))
    return False


def image_source_errors(doc, evidence):
    path = evidence.get("image_sources_file")
    if not isinstance(path, str) or not path.strip():
        return ["Image source record file missing; keep origins outside audience captions."]
    try:
        data = json.loads(Path(path).read_text(encoding="utf-8-sig"))
        records = data["images"]
        if not isinstance(records, list):
            raise ValueError("images must be a list")
        sources = {}
        for record in records:
            shape_id, url = str(record["shape_id"]), record["source_url"]
            if shape_id in sources or not isinstance(url, str) or not re.match(r"https?://[^/\s]+", url):
                raise ValueError("image IDs must be unique and origins must be HTTP(S) URLs")
            sources[shape_id] = url
    except (OSError, ValueError, KeyError, TypeError) as exc:
        return [f"Cannot validate image source record: {exc}"]
    return [f"Image {shape._inline.docPr.get('id')} is missing an original source record."
            for shape in doc.inline_shapes
            if shape._inline.docPr.get("id") not in sources]


def audit(paths, evidence, rules):
    errors, warnings, docs = [], [], {}
    if len(set(path.resolve() for path in paths.values())) != 3:
        errors.append("Three distinct DOCX files are required.")
    for role, path in paths.items():
        try:
            if path.suffix.lower() != ".docx":
                raise ValueError("Not a DOCX")
            docs[role] = Document(path)
        except Exception as exc:
            errors.append(f"Cannot read {role}: {exc}")
    if len(docs) != 3:
        return {"passed": False, "errors": errors, "warnings": warnings}

    english_text = visible_text(docs["english"])
    bilingual_text = visible_text(docs["bilingual"])
    final_doc = docs["final"]
    final_text = visible_text(final_doc)
    text_hash = hashlib.sha256(final_text.encode("utf-8")).hexdigest()
    if not re.search(r"[A-Za-z]{3}", english_text):
        errors.append("English manuscript has no recognizable English text.")
    if evidence.get("timestamps_available", True) and not TIMESTAMP.search(english_text):
        errors.append("English manuscript is missing available timestamps.")
    if TIMESTAMP.search(final_text):
        errors.append("Chinese final contains a timestamp; inspect its context.")
    if not re.search(r"[\u3400-\u9fff]", bilingual_text) or not re.search(r"[A-Za-z]{3}", bilingual_text):
        errors.append("Bilingual manuscript must contain both languages.")

    pairs = [p.style.name for p in docs["bilingual"].paragraphs
             if p.text.strip() and p.style.name in {"EnglishSentence", "ChineseTranslation"}]
    if pairs:
        if len(pairs) % 2 or any(name != ("EnglishSentence" if i % 2 == 0 else "ChineseTranslation")
                                 for i, name in enumerate(pairs)):
            errors.append("Bilingual sentence paragraph styles do not alternate correctly.")
    else:
        warnings.append("Bilingual semantic styles absent: manually inspect every sentence pair.")
    for field in ("alignment_verified", "content_review_verified"):
        if evidence.get(field) is not True:
            errors.append(f"Missing completed editorial check: {field}.")

    for stage in rules.get("experience_stages", []):
        lesson_path = SKILL / "references" / "experience" / f"{stage}.md"
        has_lessons = lesson_path.is_file() and bool(re.search(
            r"^### [A-Z]+\d+\b", lesson_path.read_text(encoding="utf-8"), re.MULTILINE))
        check = evidence.get("experience_checks", {}).get(stage, {})
        expected_status = "passed" if has_lessons else "no_experience"
        if check.get("status") != expected_status:
            errors.append(f"Stage experience self-check incomplete: {stage}.")
        if has_lessons and not check.get("checked_ids"):
            errors.append(f"Stage experience lesson IDs missing: {stage}.")
        record = check.get("record_file", "")
        if not record or not Path(record).is_file() or not Path(record).stat().st_size:
            errors.append(f"Stage experience self-check record missing: {stage}.")

    paragraphs = [p for p in final_doc.paragraphs if p.text.strip()]
    if not paragraphs or not paragraphs[0].text.startswith(rules["title_prefix"]):
        errors.append("Final title is missing the required prefix.")
    hl_index = next((i for i, p in enumerate(paragraphs) if p.text.strip() == "Z Highlights"), None)
    if hl_index is None:
        errors.append("Z Highlights heading missing.")
    before_hl = paragraphs[:hl_index] if hl_index is not None else []
    summaries = [p.text for p in before_hl if p.text.startswith("\u201c") and p.text.endswith("\u201d")]
    if not summaries or len(summaries[0]) > rules["summary_max_characters"]:
        errors.append("Missing quoted summary or summary exceeds configured limit.")

    footer = render_footer(rules)
    if "{compiler}" in final_text:
        errors.append("Unresolved compiler placeholder remains in the final article.")
    for required in [part.strip() for part in footer.split("\n\n") if part.strip()]:
        if required not in final_text:
            errors.append("Required compiler credit or full disclaimer is missing or changed.")

    counts = image_roles(final_doc)
    if counts["cover"] < rules["cover_images_minimum"]:
        errors.append("Required tagged cover image missing.")
    if counts["body"] < rules["body_images_minimum"]:
        errors.append(f"Body image minimum not met: {counts['body']}.")
    if counts["unclassified"]:
        warnings.append("Classify image alt text as ZP_COVER or ZP_BODY; cover cannot count as body.")
    errors.extend(image_source_errors(final_doc, evidence))
    if re.search(r"(?:图片来源|图源)\s*[：:]", final_text):
        errors.append("Visible image source lines must be moved to project image records.")
    for i, paragraph in enumerate(final_doc.paragraphs):
        if paragraph._p.xpath(".//w:drawing"):
            if not has_descriptive_caption(final_doc.paragraphs, i):
                errors.append(f"Image paragraph {i + 1} needs an adjacent plain description without source links or bold.")

    relationships = [rel.target_ref for rel in final_doc.part.rels.values()
                     if rel.is_external and rel.reltype.endswith("/hyperlink")]
    if not any(str(url).startswith(("https://", "http://")) for url in relationships) \
            and not re.search(r"https?://", final_text):
        errors.append("Original/source links missing.")

    spaces = re.findall(r"[\u3400-\u9fff] +[A-Za-z0-9]|[A-Za-z0-9%] +[\u3400-\u9fff]", final_text)
    if spaces:
        warnings.append(f"Inspect {len(spaces)} Chinese/English/number boundary spaces; do not globally delete.")

    comments, anchors = comment_map(paths["final"])
    review = evidence.get("doubao", {})
    status = review.get("status", "pending")
    if status not in {"no_risk", "risk_found"}:
        errors.append("Actual Doubao full-text review is incomplete.")
    else:
        if review.get("prompt") != rules["review_prompt"]:
            errors.append("Doubao prompt differs from the requested prompt.")
        if review.get("text_sha256") != text_hash:
            errors.append("Doubao reviewed text differs from the current final text.")
        try:
            datetime.fromisoformat(review.get("reviewed_at", "").replace("Z", "+00:00"))
        except (ValueError, TypeError):
            errors.append("Actual Doubao review date/time missing or invalid.")
        responses = review.get("response_files", [])
        if not responses or not all(Path(p).is_file() and Path(p).stat().st_size > 0 for p in responses):
            errors.append("Saved actual Doubao responses missing.")
        if status == "no_risk" and not any(
                rules["no_risk_comment"] in text and text_hash in text and "已失效" not in text
                and cid in anchors for cid, text in comments.items()):
            errors.append("Anchored full-text no-risk review comment missing.")
        if status == "risk_found":
            ids = [str(cid) for cid in review.get("comment_ids", [])]
            if not ids or not all(cid in comments and cid in anchors for cid in ids):
                errors.append("Risk findings are not mapped to anchored Word comments.")

    for role, path in paths.items():
        item = evidence.get("visual_review", {}).get(role, {})
        if item.get("reviewed") is not True or item.get("docx_sha256") != sha256(path):
            errors.append(f"Missing or stale visual review for {role}.")
        pages = item.get("page_images", [])
        if not pages or not all(Path(p).is_file() and Path(p).stat().st_size > 0 for p in pages):
            errors.append(f"Rendered page evidence missing for {role}.")

    return {"passed": not errors, "errors": errors, "warnings": warnings,
            "image_counts": counts, "final_text_sha256": text_hash,
            "comment_count": len(comments),
            "note": "Static checks only; authenticity, translation, content and every-page appearance require actual review."}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    for role in ("english", "bilingual", "final"):
        parser.add_argument(f"--{role}", required=True, type=Path)
    parser.add_argument("--evidence", type=Path)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--extract-final", type=Path)
    parser.add_argument("--profile", type=Path)
    parser.add_argument("--compiler")
    args = parser.parse_args()
    try:
        rules = load_rules(args.profile, args.compiler)
    except (OSError, ValueError) as exc:
        parser.error(str(exc))
    evidence = json.loads(args.evidence.read_text(encoding="utf-8")) if args.evidence else {}
    paths = {role: getattr(args, role) for role in ("english", "bilingual", "final")}
    if args.extract_final:
        text = visible_text(Document(paths["final"]))
        args.extract_final.parent.mkdir(parents=True, exist_ok=True)
        args.extract_final.write_text(text, encoding="utf-8")
    result = audit(paths, evidence, rules)
    serialized = json.dumps(result, ensure_ascii=False, indent=2)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(serialized + "\n", encoding="utf-8")
    print(serialized)
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    sys.exit(main())

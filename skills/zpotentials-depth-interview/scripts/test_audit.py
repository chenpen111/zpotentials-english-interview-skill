"""Synthetic structural tests only: no real article or Doubao review is performed."""

import importlib.util
import json
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches
from PIL import Image


HERE = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location("audit_delivery", HERE / "audit_delivery.py")
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


def run_tests():
    caption_fixture = Document()
    caption_fixture.add_paragraph("IMAGE PLACEHOLDER")
    caption_fixture.add_paragraph("图片说明。")
    assert module.has_descriptive_caption(caption_fixture.paragraphs, 0)
    caption_fixture.paragraphs[1].runs[0].bold = True
    assert not module.has_descriptive_caption(caption_fixture.paragraphs, 0)
    caption_fixture.paragraphs[1].runs[0].bold = False
    caption_fixture.paragraphs[1].text = "图片来源：TEST FIXTURE"
    assert not module.has_descriptive_caption(caption_fixture.paragraphs, 0)
    caption_fixture.paragraphs[1].text = "图片说明 https://example.com"
    assert not module.has_descriptive_caption(caption_fixture.paragraphs, 0)
    caption_fixture.paragraphs[1].text = "另一主题"
    caption_fixture.paragraphs[1].style = "Heading 2"
    assert not module.has_descriptive_caption(caption_fixture.paragraphs, 0)
    assert not module.has_descriptive_caption(caption_fixture.paragraphs, 1)
    rules = module.load_rules(compiler="Test Compiler", use_local=False)
    footer = module.render_footer(rules)
    assert "编译：Test Compiler" in footer and "{compiler}" not in footer
    try:
        module.load_rules(use_local=False)
    except ValueError:
        pass
    else:
        raise AssertionError("A new colleague must not silently inherit another compiler's credit")
    with TemporaryDirectory(prefix="zp-skill-test-") as temp:
        root = Path(temp)
        profile = root / "test-profile.json"
        profile.write_text(json.dumps({"compiler": "Profile Author"}), encoding="utf-8")
        assert module.load_rules(profile, use_local=False)["compiler"] == "Profile Author"
        assert module.load_rules(profile, "Task Author", use_local=False)["compiler"] == "Task Author"
        profile.write_text(json.dumps({"compiler": "Author", "unknown": True}), encoding="utf-8")
        try:
            module.load_rules(profile, use_local=False)
        except ValueError:
            pass
        else:
            raise AssertionError("Unknown profile keys must not be silently accepted")
        image = root / "synthetic.png"
        Image.new("RGB", (100, 60), "gray").save(image)
        paths = {role: root / f"{role}.docx" for role in ("english", "bilingual", "final")}
        english = Document()
        english.add_paragraph("TEST FIXTURE Speaker 0:00 Hello. Thank you. 0:10 Goodbye.")
        english.save(paths["english"])
        bilingual = Document()
        for name in ("EnglishSentence", "ChineseTranslation"):
            bilingual.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        bilingual.add_paragraph("TEST FIXTURE Speaker: Hello.", style="EnglishSentence")
        bilingual.add_paragraph("Speaker：你好。", style="ChineseTranslation")
        bilingual.save(paths["bilingual"])
        final = Document()
        final.add_paragraph(rules["title_prefix"] + "TEST FIXTURE", style="Title")
        author = final.add_paragraph("作者：Test Interview Publisher")
        summary = final.add_paragraph("摘要：“这是测试，不是生产稿件。”")
        cover = final.add_paragraph().add_run().add_picture(str(image), width=Inches(1))
        cover._inline.docPr.set("descr", "ZP_COVER")
        final.add_paragraph("测试封面：用于演示排版的灰色矩形。")
        final.add_paragraph("Z Highlights", style="Heading 2")
        highlight = final.add_paragraph("这是模拟内容。")
        final.add_paragraph("这也是模拟内容。")
        final.add_paragraph("第三条模拟亮点。")
        final.add_paragraph().add_run("这是访谈背景介绍。test fixture only").italic = True
        image_sources = [{"shape_id": cover._inline.docPr.get("id"),
                          "source_url": "https://example.com/test-fixture"}]
        for i in range(4):
            shape = final.add_paragraph().add_run().add_picture(str(image), width=Inches(1))
            shape._inline.docPr.set("descr", "ZP_BODY")
            final.add_paragraph("测试图片：用于演示排版的灰色矩形。")
            image_sources.append({"shape_id": shape._inline.docPr.get("id"),
                                  "source_url": "https://example.com/test-fixture"})
        source_file = root / "image-sources.json"
        source_file.write_text(json.dumps({"images": image_sources}), encoding="utf-8")
        title = final.add_paragraph("原视频：Test Interview Original Title")
        url = final.add_paragraph("https://example.com/test-interview")
        for part in footer.strip().split("\n\n"):
            final.add_paragraph(part)
        text_hash = module.hashlib.sha256(module.visible_text(final).encode("utf-8")).hexdigest()
        final.add_comment(final.paragraphs[-1].runs, rules["no_risk_comment"] + "；TEST FIXTURE ONLY；" + text_hash, author="Test")
        final.save(paths["final"])
        response = root / "synthetic-review.txt"
        response.write_text("TEST FIXTURE, NOT AN ACTUAL DOUBAO RESPONSE", encoding="utf-8")
        self_checks = root / "synthetic-self-checks.md"
        self_checks.write_text("TEST FIXTURE ONLY: synthetic stage self-check records.", encoding="utf-8")
        evidence = {
            "alignment_verified": True, "content_review_verified": True,
            "interview_credit": "Test Interview Publisher",
            "original_title": "Test Interview Original Title",
            "original_url": "https://example.com/test-interview",
            "image_sources_file": str(source_file),
            "experience_checks": {stage: {
                "status": "passed", "record_file": str(self_checks),
                "checked_ids": module.re.findall(r"^### ([A-Z]+\d+)\b", (
                    module.SKILL / "references" / "experience" / f"{stage}.md"
                ).read_text(encoding="utf-8"), module.re.MULTILINE)
            } for stage in rules["experience_stages"]},
            "doubao": {
                "status": "no_risk", "reviewed_at": "2026-10-14T12:00:00+08:00",
                "prompt": rules["review_prompt"],
                "text_sha256": module.hashlib.sha256(module.visible_text(final).encode("utf-8")).hexdigest(),
                "response_files": [str(response)]
            },
            "visual_review": {role: {"reviewed": True, "docx_sha256": module.sha256(path),
                                     "page_images": [str(image)]} for role, path in paths.items()}
        }
        result = module.audit(paths, evidence, rules)
        assert result["passed"], result

        summary_text = summary.text
        for invalid in ["“缺少摘要标签。”", "摘要：“内容。”————Test Speaker",
                        "摘要：内容。", "摘要：内容。————Test Compiler",
                        "摘要：“”", "摘要：“" + "长" * 81 + "”"]:
            summary.text = invalid
            assert module.packaging_errors(final, evidence, rules)
        summary.text = summary_text
        for invalid in ["“亮点无需引号。”", "> 不用引用块", "内容[1]", "内容https://example.com"]:
            highlight.text = invalid
            assert module.packaging_errors(final, evidence, rules)
        highlight.text = "这是模拟内容。"
        highlight.style = "Quote"
        assert module.packaging_errors(final, evidence, rules)
        highlight.style = "Normal"
        opening_credit = summary.insert_paragraph_before("原访谈：TEST FIXTURE")
        assert module.packaging_errors(final, evidence, rules)
        opening_credit._p.getparent().remove(opening_credit._p)
        author.text = "作者：Test Guest Not Publisher"
        assert module.packaging_errors(final, evidence, rules)
        author.text = "作者：" + evidence["interview_credit"]
        summary._p.addnext(author._p)
        assert module.packaging_errors(final, evidence, rules)
        summary._p.addprevious(author._p)
        cover_paragraph = cover._inline.xpath("ancestor::w:p")[0]
        cover_paragraph.addnext(summary._p)
        assert module.packaging_errors(final, evidence, rules)
        cover_paragraph.addprevious(summary._p)
        title.text = "原视频：Incorrect Title"
        assert module.packaging_errors(final, evidence, rules)
        title.text = "原视频：" + evidence["original_title"]
        for invalid in ["原始链接：" + evidence["original_url"],
                        "[观看视频](" + evidence["original_url"] + ")",
                        "https://example.com/wrong-video"]:
            url.text = invalid
            assert module.packaging_errors(final, evidence, rules)
        url.text = evidence["original_url"]
        from docx.oxml import OxmlElement
        field = OxmlElement("w:fldSimple")
        url._p.append(field)
        assert module.packaging_errors(final, evidence, rules)
        url._p.remove(field)
        assert not module.packaging_errors(final, evidence, rules), module.packaging_errors(final, evidence, rules)
        final.save(paths["final"])

        missing_sources = dict(evidence)
        missing_sources.pop("image_sources_file")
        assert module.image_source_errors(final, missing_sources)
        for invalid in [[], [image_sources[0], image_sources[0]],
                        [{"shape_id": "1", "source_url": "not-a-url"}]]:
            source_file.write_text(json.dumps({"images": invalid}), encoding="utf-8")
            assert module.image_source_errors(final, evidence)
        source_file.write_text(json.dumps({"images": image_sources}), encoding="utf-8")
        assert not module.image_source_errors(final, evidence)

        final.add_paragraph("图片来源：https://example.com/old-caption")
        final.save(paths["final"])
        assert any("Visible image source" in msg for msg in module.audit(paths, evidence, rules)["errors"])
        old_caption = final.paragraphs[-1]._p
        old_caption.getparent().remove(old_caption)
        final.save(paths["final"])

        unchecked = json.loads(json.dumps(evidence))
        unchecked["experience_checks"]["translation"]["status"] = "needs_fix"
        assert any("Stage experience" in msg for msg in module.audit(paths, unchecked, rules)["errors"])
        del unchecked["experience_checks"]["images"]
        assert any("images" in msg for msg in module.audit(paths, unchecked, rules)["errors"])

        pending = json.loads(json.dumps(evidence))
        pending["doubao"]["status"] = "pending"
        assert any("incomplete" in msg for msg in module.audit(paths, pending, rules)["errors"])

        first_comment = next(iter(final.comments))
        first_comment.add_paragraph("已失效：TEST FIXTURE")
        final.save(paths["final"])
        assert any("no-risk" in msg for msg in module.audit(paths, evidence, rules)["errors"])
        expired = first_comment.paragraphs[-1]._p
        expired.getparent().remove(expired)

        final.add_paragraph("Speaker：0:12 This added text makes the review stale.")
        final.save(paths["final"])
        failed = module.audit(paths, evidence, rules)
        assert any("timestamp" in msg for msg in failed["errors"])
        assert any("differs" in msg for msg in failed["errors"])
        assert any("stale" in msg for msg in failed["errors"])

        missing = Document(paths["final"])
        missing.paragraphs[-2]._p.getparent().remove(missing.paragraphs[-2]._p)
        missing.inline_shapes[-1]._inline.getparent().remove(missing.inline_shapes[-1]._inline)
        missing.save(paths["final"])
        failed = module.audit(paths, evidence, rules)
        assert any("disclaimer" in msg for msg in failed["errors"])
        assert any("minimum" in msg for msg in failed["errors"])

        bilingual.paragraphs[-1].style = "EnglishSentence"
        bilingual.save(paths["bilingual"])
        assert any("alternate" in msg for msg in module.audit(paths, evidence, rules)["errors"])

        aliases = dict(paths)
        aliases["bilingual"] = paths["english"]
        assert any("distinct" in msg for msg in module.audit(aliases, evidence, rules)["errors"])
        print("PASS: verified interview credit and quoted summary before first image; plain Highlights; original-video title then plain URL; captions and image sources; stage checks; review state; stale text/render; timestamps; footer/images; bilingual pairs; distinct files.")
        print("Synthetic tests only. No real Doubao access, rendering or article revision performed.")


if __name__ == "__main__":
    run_tests()
